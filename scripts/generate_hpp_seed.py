"""Generate Genresi stock seed XML from the HPP workbook."""

# pylint: disable=import-error,too-many-locals

from __future__ import annotations

import argparse
import warnings
from collections.abc import Iterable
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl

DEFAULT_WORKBOOK = Path("docs/HPP_ROTI_REVISI__20_.xlsx")
DEFAULT_OUTPUT = Path("addons/genresi_stock/data")
UOM_REFS = {
    "gram": "uom.product_uom_gram",
    "g": "uom.product_uom_gram",
    "pcs": "genresi_base.uom_genresi_pcs",
    "biji": "genresi_base.uom_genresi_biji",
}
MIN_STOCK = {
    "Ragi permifan 2 in 1": 500,
    "Telor": 2000,
    "Terigu Golden eagle": 5000,
    "Whole wheat natural": 4000,
    "Gula pasir": 2000,
    "Mother choice": 1000,
    "Garam": 500,
    "Keju": 800,
    "Coklat ceres": 1000,
    "Daging": 1500,
    "Elmer glaze": 1000,
    "SKM": 500,
    "Selai srikaya": 500,
    "Lunch box": 100,
    "Kresek": 100,
}
MATERIAL_IDS = {
    "Terigu Golden eagle": "material_terigu_golden_eagle",
    "Terigu gandum whole wheat bungasari": "material_whole_wheat_bungasari",
    "Gandum Kasar": "material_gandum_kasar",
    "Whole wheat natural": "material_whole_wheat_natural",
    "Ragi permifan 2 in 1": "material_ragi",
    "Gula pasir": "material_gula",
    "Telor": "material_telor",
    "Mother choice": "material_mother_choice",
    "Garam": "material_garam",
    "Coklat": "material_coklat",
    "Keju": "material_keju",
    "Elmer glaze": "material_elmer_glaze",
    "SKM": "material_skm",
    "Coklat ceres": "material_coklat_ceres",
    "Daging": "material_daging",
    "Lunch box": "material_lunch_box",
    "Kresek": "material_kresek",
    "Kundur": "material_kundur",
    "Bombay": "material_bombay",
}
LINE_ID_PARTS = {
    "Terigu gandum whole wheat bungasari": "flour",
    "Whole wheat natural": "flour",
    "Terigu Golden eagle": "terigu",
    "Ragi permifan 2 in 1": "ragi",
    "Gula pasir": "gula",
    "Telor": "telor",
    "Mother choice": "mother",
    "Garam": "garam",
    "Elmer glaze": "glaze",
    "Coklat ceres": "",
    "Lunch box": "lunch_box",
    "Kresek": "kresek",
    "Selai srikaya": "",
}
PRODUCT_IDS = {
    "Roti Original - Adonan Bungasari": "product_original_bungasari",
    "Roti Original - Adonan Natural": "product_original_natural",
    "Roti Keju - Adonan Bungasari": "product_keju_bungasari",
    "Roti Keju - Adonan Natural": "product_keju_natural",
    "Roti Daging - Adonan Bungasari": "product_daging_bungasari",
    "Roti Daging - Adonan Natural": "product_daging_natural",
    "Roti Coklat - Adonan Bungasari": "product_coklat_bungasari",
    "Roti Coklat - Adonan Natural": "product_coklat_natural",
}
PRODUCT_PRICES = {
    "Roti Original - Adonan Bungasari": 20000,
    "Roti Original - Adonan Natural": 20000,
    "Roti Keju - Adonan Bungasari": 25000,
    "Roti Keju - Adonan Natural": 25000,
    "Roti Daging - Adonan Bungasari": 35000,
    "Roti Daging - Adonan Natural": 35000,
    "Roti Coklat - Adonan Bungasari": 25000,
    "Roti Coklat - Adonan Natural": 25000,
}
RECIPE_DEFINITIONS = (
    ("recipe_dough_bungasari", "Adonan Bungasari", "dough", 6, 12, 14, 0),
    ("recipe_dough_natural", "Adonan Natural", "dough", 19, 25, 27, 0),
    ("recipe_filling_keju", "Filling Keju", "filling", 32, 34, 36, 0),
    ("recipe_filling_daging", "Filling Daging", "filling", 41, 43, 47, 45),
    ("recipe_filling_coklat", "Filling Coklat", "filling", 52, 52, 54, 0),
    ("recipe_packaging", "Kemasan", "packaging", 59, 60, None, 0),
)
EXTRA_MATERIALS = {
    "Selai srikaya": {
        "id": "material_selai_srikaya",
        "uom_ref": "uom.product_uom_gram",
        "price_per_unit": 0,
        "pack_size": 1,
    }
}


def _number(value) -> str:
    """Format workbook numbers without unnecessary trailing zeroes."""
    if value is None:
        return "0"
    number = Decimal(str(value))
    if number == number.to_integral():
        return str(int(number))
    return format(number.normalize(), "f")


def _slug(value: str) -> str:
    """Create a stable lowercase identifier fragment."""
    return "".join(character.lower() if character.isalnum() else "_" for character in value).strip(
        "_"
    )


def _field(record, name: str, value: str | int | float | None = None, ref: str | None = None):
    """Append one Odoo XML field to a record."""
    attributes = {"name": name}
    if ref:
        attributes["ref"] = ref
    element = ET.SubElement(record, "field", attributes)
    if value is not None:
        element.text = str(value)
    return element


def _record(root, record_id: str, model: str, fields: Iterable[tuple[str, object, str | None]]):
    """Append one Odoo XML record."""
    record = ET.SubElement(root, "record", {"id": record_id, "model": model})
    for name, value, ref in fields:
        _field(record, name, value, ref)
    return record


def _write_xml(path: Path, root: ET.Element):
    """Write an indented Odoo XML document."""
    ET.indent(root, space="    ")
    tree = ET.ElementTree(root)
    tree.write(path, encoding="utf-8", xml_declaration=True)


def _read_materials(sheet):
    """Read material prices and package sizes from the price sheet."""
    materials = {}
    for row in sheet.iter_rows(min_row=4, max_row=22, values_only=True):
        number, name, purchase_price, pack_size, unit, price_per_unit = row[:6]
        if not isinstance(number, int) or not name:
            continue
        unit_key = str(unit).strip().lower()
        if unit_key not in UOM_REFS:
            raise ValueError(f"Satuan workbook tidak didukung: {unit}")
        materials[name] = {
            "id": MATERIAL_IDS.get(name, f"material_{_slug(name)}"),
            "uom_ref": UOM_REFS[unit_key],
            "price_per_unit": price_per_unit,
            "pack_size": pack_size,
            "purchase_price": purchase_price,
            "min_stock": MIN_STOCK.get(name, 0),
        }
    for name, values in EXTRA_MATERIALS.items():
        materials[name] = values | {"min_stock": MIN_STOCK.get(name, 0)}
    return materials


def _read_recipes(sheet, materials):
    """Read recipe sections and their batch outputs from the HPP sheet."""
    recipes = []
    for (
        recipe_id,
        name,
        recipe_type,
        first_row,
        last_row,
        batch_row,
        additional_row,
    ) in RECIPE_DEFINITIONS:
        lines = []
        for row in sheet.iter_rows(min_row=first_row, max_row=last_row, values_only=True):
            item, quantity, unit = row[:3]
            if not item or not isinstance(quantity, int | float):
                continue
            if item not in materials:
                raise ValueError(f"Bahan workbook tidak ditemukan: {item}")
            lines.append((item, quantity, unit))
        batch_output = 1 if batch_row is None else sheet.cell(batch_row, 5).value
        additional_cost = 0
        if additional_row:
            additional_cost = sheet.cell(additional_row, 5).value or 0
        recipes.append(
            {
                "id": recipe_id,
                "name": name,
                "recipe_type": recipe_type,
                "batch_output_qty": batch_output,
                "additional_cost": additional_cost,
                "lines": lines,
                "active": recipe_id != "recipe_dough_bungasari",
                "state": "active" if recipe_id != "recipe_dough_bungasari" else "archived",
            }
        )

    recipes.append(
        {
            "id": "recipe_filling_srikaya",
            "name": "Filling Srikaya",
            "recipe_type": "filling",
            "batch_output_qty": 15,
            "additional_cost": 0,
            "lines": [("Selai srikaya", 1, "gram")],
            "active": True,
            "state": "active",
        }
    )
    return recipes


def _read_products(summary_sheet, margin_sheet, recipes):
    """Read product names and sale prices from the workbook summaries."""
    recipe_ids = {recipe["name"]: recipe["id"] for recipe in recipes}
    products = []
    for index, row in enumerate(
        summary_sheet.iter_rows(min_row=65, max_row=72, values_only=True), start=6
    ):
        name = row[0]
        if not name:
            continue
        dough_name = "Adonan Bungasari" if "Bungasari" in name else "Adonan Natural"
        filling_name = next(
            (
                f"Filling {filling}"
                for filling in ("Keju", "Daging", "Coklat")
                if f" {filling} " in name
            ),
            None,
        )
        sale_price = margin_sheet.cell(index, 3).value or PRODUCT_PRICES.get(name, 0)
        products.append(
            {
                "id": PRODUCT_IDS.get(name, f"product_{_slug(name)}"),
                "name": name.replace("Roti Coklat", "Roti Cokelat"),
                "dough_recipe_id": recipe_ids[dough_name],
                "filling_recipe_id": recipe_ids.get(filling_name),
                "packaging_recipe_id": recipe_ids["Kemasan"],
                "sale_price": sale_price,
                "active": dough_name != "Adonan Bungasari",
                "state": "active" if dough_name != "Adonan Bungasari" else "archived",
            }
        )
    products.append(
        {
            "id": "product_srikaya_natural",
            "name": "Roti Srikaya - Adonan Natural",
            "dough_recipe_id": recipe_ids["Adonan Natural"],
            "filling_recipe_id": recipe_ids["Filling Srikaya"],
            "packaging_recipe_id": recipe_ids["Kemasan"],
            "sale_price": 25000,
            "active": True,
            "state": "active",
        }
    )
    return products


def _materials_xml(materials):
    root = ET.Element("odoo")
    for name, material in materials.items():
        values = [
            ("name", name, None),
            ("uom_id", None, material["uom_ref"]),
            ("price_per_unit", _number(material["price_per_unit"]), None),
            ("pack_size", _number(material["pack_size"]), None),
        ]
        if material.get("min_stock"):
            values.append(("min_stock", _number(material["min_stock"]), None))
        _record(root, material["id"], "genresi.material", values)
    return root


def _recipes_xml(recipes, materials):
    root = ET.Element("odoo")
    for recipe in recipes:
        values = [
            ("name", recipe["name"], None),
            ("recipe_type", recipe["recipe_type"], None),
            ("batch_output_qty", _number(recipe["batch_output_qty"]), None),
        ]
        if recipe["additional_cost"]:
            values.append(("additional_cost", _number(recipe["additional_cost"]), None))
        if not recipe["active"]:
            values.extend([("active", "False", None), ("state", recipe["state"], None)])
        _record(root, recipe["id"], "genresi.recipe", values)
        for item, quantity, _unit in recipe["lines"]:
            line_part = LINE_ID_PARTS.get(item, _slug(item))
            line_id = f"{recipe['id']}_line{('_' + line_part) if line_part else ''}"
            _record(
                root,
                line_id,
                "genresi.recipe.line",
                [
                    ("recipe_id", None, recipe["id"]),
                    ("material_id", None, materials[item]["id"]),
                    ("quantity", _number(quantity), None),
                ],
            )
    return root


def _products_xml(products):
    root = ET.Element("odoo")
    for product in products:
        values = [
            ("name", product["name"], None),
            ("dough_recipe_id", None, product["dough_recipe_id"]),
        ]
        if product["filling_recipe_id"]:
            values.append(("filling_recipe_id", None, product["filling_recipe_id"]))
        values.extend(
            [
                ("packaging_recipe_id", None, product["packaging_recipe_id"]),
                ("sale_price", _number(product["sale_price"]), None),
            ]
        )
        if not product["active"]:
            values.extend([("active", "False", None), ("state", product["state"], None)])
        _record(root, product["id"], "genresi.product", values)
    return root


def generate(workbook_path: Path, output_dir: Path):
    """Generate the three stock seed XML files and return their counts."""
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Conditional Formatting extension")
        workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    required_sheets = {"Harga Bahan", "Resep & HPP", "Margin & Simulasi Sewa"}
    missing = required_sheets.difference(workbook.sheetnames)
    if missing:
        raise ValueError(f"Sheet workbook tidak ditemukan: {', '.join(sorted(missing))}")

    materials = _read_materials(workbook["Harga Bahan"])
    recipes = _read_recipes(workbook["Resep & HPP"], materials)
    products = _read_products(workbook["Resep & HPP"], workbook["Margin & Simulasi Sewa"], recipes)
    output_dir.mkdir(parents=True, exist_ok=True)
    _write_xml(output_dir / "materials.xml", _materials_xml(materials))
    _write_xml(output_dir / "recipes.xml", _recipes_xml(recipes, materials))
    _write_xml(output_dir / "products.xml", _products_xml(products))
    return len(materials), len(recipes), len(products)


def main():
    """Parse arguments, generate seeds, and print a compact summary."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    counts = generate(args.workbook, args.output_dir)
    print(
        f"Generated HPP seed: materials={counts[0]}, recipes={counts[1]}, "
        f"products={counts[2]} -> {args.output_dir}"
    )


if __name__ == "__main__":
    main()
